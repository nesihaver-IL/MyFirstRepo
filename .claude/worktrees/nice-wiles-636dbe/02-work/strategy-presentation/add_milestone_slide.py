#!/usr/bin/env python3
"""
Add a Major Milestones slide to Strategy_active_format.pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def add_milestone_slide(prs, title="Major Milestones", content="Project milestones and key deliverables will be detailed here."):
    """Add a milestones slide to the presentation."""
    layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(layout)

    # Add header bar (matching corporate color scheme)
    header = slide.shapes.add_shape(
        1,  # Rectangle
        Inches(0), Inches(0), prs.slide_width, Inches(1.0)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(0x1F, 0x4E, 0x79)  # Dark navy
    header.line.fill.background()

    # Add title text
    tf = header.text_frame
    tf.text = title
    tf.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    tf.paragraphs[0].runs[0].font.bold = True
    tf.paragraphs[0].runs[0].font.size = Pt(28)
    tf.paragraphs[0].alignment = PP_ALIGN.LEFT
    tf.margin_left = Inches(0.5)
    tf.margin_top = Inches(0.15)

    # Add content placeholder
    content_box = slide.shapes.add_textbox(
        Inches(0.5), Inches(1.5), Inches(12.33), Inches(5.5)
    )
    tf = content_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = content
    p.runs[0].font.size = Pt(16)
    p.runs[0].font.color.rgb = RGBColor(0x70, 0x70, 0x70)
    p.runs[0].font.italic = True

    print(f"✓ Added '{title}' slide to presentation")

def main():
    pptx_file = "/home/nhaver/MyFirstRepo/03-plans/strategy-presentation/Strategy_active_format.pptx"

    # Load existing presentation
    prs = Presentation(pptx_file)
    print(f"Loaded presentation with {len(prs.slides)} slides")

    # Add milestone slide
    add_milestone_slide(prs)

    # Save updated presentation
    prs.save(pptx_file)
    print(f"✓ Saved updated presentation: {pptx_file}")
    print(f"  Total slides: {len(prs.slides)}")

if __name__ == "__main__":
    main()
