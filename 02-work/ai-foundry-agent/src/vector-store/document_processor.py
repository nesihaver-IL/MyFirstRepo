"""Document processor for preparing documents for vector storage."""

import hashlib
import logging
from typing import List, Dict, Any, Optional
from pathlib import Path
import json
from pptx import Presentation
from PIL import Image
import io

logger = logging.getLogger(__name__)


class DocumentProcessor:
    """Processes documents for ingestion into vector store."""

    def __init__(
        self,
        chunk_size: int = 1000,
        chunk_overlap: int = 200,
        image_describer=None,
    ):
        """Initialize Document Processor.

        Args:
            chunk_size: Maximum size of text chunks in characters
            chunk_overlap: Number of characters to overlap between chunks
            image_describer: Optional ImageDescriber instance for processing images
        """
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.image_describer = image_describer

    def chunk_text(self, text: str) -> List[str]:
        """Split text into overlapping chunks.

        Args:
            text: Text to chunk

        Returns:
            List of text chunks
        """
        if len(text) <= self.chunk_size:
            return [text]

        chunks = []
        start = 0

        while start < len(text):
            end = start + self.chunk_size

            # Try to break at sentence boundary
            if end < len(text):
                # Look for sentence endings
                for punct in [". ", ".\n", "! ", "?\n", "? "]:
                    last_punct = text.rfind(punct, start, end)
                    if last_punct != -1:
                        end = last_punct + len(punct)
                        break

            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)

            start = end - self.chunk_overlap

        return chunks

    def process_document(
        self,
        content: str,
        title: str,
        category: str = "general",
        metadata: Optional[Dict[str, Any]] = None,
        source_id: Optional[str] = None,
    ) -> List[Dict[str, Any]]:
        """Process a document into searchable chunks.

        Args:
            content: Document content
            title: Document title
            category: Document category
            metadata: Additional metadata
            source_id: Source document identifier

        Returns:
            List of processed document chunks ready for upload
        """
        # Chunk the content
        chunks = self.chunk_text(content)

        # Create document records
        documents = []
        for i, chunk in enumerate(chunks):
            # Generate unique ID
            chunk_id = self._generate_id(source_id or title, i)

            # Prepare metadata
            chunk_metadata = {
                "source_id": source_id or title,
                "chunk_index": i,
                "total_chunks": len(chunks),
                **(metadata or {}),
            }

            # Create document record
            doc = {
                "id": chunk_id,
                "content": chunk,
                "title": f"{title} (Part {i + 1}/{len(chunks)})"
                if len(chunks) > 1
                else title,
                "category": category,
                "metadata": json.dumps(chunk_metadata),
            }

            documents.append(doc)

        logger.info(
            f"Processed document '{title}' into {len(documents)} chunks"
        )
        return documents

    def process_file(
        self,
        file_path: str,
        category: str = "general",
        metadata: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """Process a file into searchable chunks.

        Args:
            file_path: Path to the file
            category: Document category
            metadata: Additional metadata

        Returns:
            List of processed document chunks
        """
        path = Path(file_path)

        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # Determine file type and route to appropriate handler
        suffix = path.suffix.lower()

        # PowerPoint files
        if suffix in [".ppt", ".pptx"]:
            return self._process_ppt_file(path, category, metadata)

        # Image files
        elif suffix in [".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp"]:
            return self._process_image_file(path, category, metadata)

        # Text files
        else:
            return self._process_text_file(path, category, metadata)

    def _process_text_file(
        self,
        path: Path,
        category: str = "general",
        metadata: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """Process a text file.

        Args:
            path: Path to the file
            category: Document category
            metadata: Additional metadata

        Returns:
            List of processed document chunks
        """
        # Read file content
        try:
            with open(path, "r", encoding="utf-8") as f:
                content = f.read()
        except Exception as e:
            logger.error(f"Error reading file {path}: {e}")
            raise

        # Use filename as title
        title = path.stem

        # Add file metadata
        file_metadata = {
            "filename": path.name,
            "file_extension": path.suffix,
            **(metadata or {}),
        }

        return self.process_document(
            content=content,
            title=title,
            category=category,
            metadata=file_metadata,
            source_id=str(path),
        )

    def _process_ppt_file(
        self,
        path: Path,
        category: str = "general",
        metadata: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """Process a PowerPoint file.

        Args:
            path: Path to the PPT file
            category: Document category
            metadata: Additional metadata

        Returns:
            List of processed document chunks
        """
        try:
            prs = Presentation(str(path))
            content_parts = []

            logger.info(f"Processing PowerPoint: {path.name} ({len(prs.slides)} slides)")

            for slide_num, slide in enumerate(prs.slides, 1):
                # Extract text from slide
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text.strip():
                        slide_texts.append(shape.text.strip())

                # Extract images from slide
                if self.image_describer:
                    for shape in slide.shapes:
                        if shape.shape_type == 13:  # Picture type
                            try:
                                # Extract image from shape
                                image = shape.image
                                image_bytes = image.blob

                                # Generate description
                                slide_context = (
                                    f"Slide {slide_num}"
                                    + (f": {slide_texts[0]}" if slide_texts else "")
                                )
                                description = self.image_describer.describe_image_from_bytes(
                                    image_bytes=image_bytes,
                                    image_name=f"{path.stem}_slide{slide_num}",
                                    context=slide_context,
                                )

                                slide_texts.append(f"\n[Image Description: {description}]")
                                logger.info(
                                    f"Described image on slide {slide_num} of {path.name}"
                                )
                            except Exception as e:
                                logger.warning(
                                    f"Could not process image on slide {slide_num}: {e}"
                                )

                # Combine slide content
                if slide_texts:
                    slide_content = f"\n--- Slide {slide_num} ---\n" + "\n".join(
                        slide_texts
                    )
                    content_parts.append(slide_content)

            # Combine all slides
            full_content = "\n\n".join(content_parts)

            if not full_content.strip():
                logger.warning(f"No content extracted from {path.name}")
                full_content = f"[PowerPoint file: {path.name} - no extractable content]"

            # Add file metadata
            file_metadata = {
                "filename": path.name,
                "file_extension": path.suffix,
                "slide_count": len(prs.slides),
                **(metadata or {}),
            }

            return self.process_document(
                content=full_content,
                title=path.stem,
                category=category,
                metadata=file_metadata,
                source_id=str(path),
            )

        except Exception as e:
            logger.error(f"Error processing PowerPoint file {path}: {e}")
            raise

    def _process_image_file(
        self,
        path: Path,
        category: str = "general",
        metadata: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """Process an image file.

        Args:
            path: Path to the image file
            category: Document category
            metadata: Additional metadata

        Returns:
            List of processed document chunks
        """
        if not self.image_describer:
            logger.warning(
                f"No image describer available, skipping image: {path.name}"
            )
            return []

        try:
            # Generate description
            description = self.image_describer.describe_image(
                image_path=str(path),
                context=f"Image file: {path.name}",
            )

            # Create content with image description
            content = f"[Image: {path.name}]\n\n{description}"

            # Add file metadata
            file_metadata = {
                "filename": path.name,
                "file_extension": path.suffix,
                "content_type": "image",
                **(metadata or {}),
            }

            return self.process_document(
                content=content,
                title=path.stem,
                category=category,
                metadata=file_metadata,
                source_id=str(path),
            )

        except Exception as e:
            logger.error(f"Error processing image file {path}: {e}")
            raise

    def process_directory(
        self,
        directory_path: str,
        category: str = "general",
        file_pattern: str = "*.md",
        metadata: Optional[Dict[str, Any]] = None,
    ) -> List[Dict[str, Any]]:
        """Process all files in a directory.

        Args:
            directory_path: Path to the directory
            category: Document category
            file_pattern: Glob pattern for files to process
            metadata: Additional metadata

        Returns:
            List of all processed document chunks
        """
        path = Path(directory_path)

        if not path.exists() or not path.is_dir():
            raise ValueError(f"Invalid directory: {directory_path}")

        all_documents = []
        files = list(path.glob(file_pattern))

        logger.info(
            f"Processing {len(files)} files from {directory_path}"
        )

        for file_path in files:
            try:
                docs = self.process_file(
                    str(file_path), category=category, metadata=metadata
                )
                all_documents.extend(docs)
            except Exception as e:
                logger.error(
                    f"Error processing file {file_path}: {e}"
                )
                continue

        logger.info(
            f"Processed {len(files)} files into {len(all_documents)} chunks"
        )
        return all_documents

    @staticmethod
    def _generate_id(source: str, index: int) -> str:
        """Generate a unique ID for a document chunk.

        Args:
            source: Source identifier
            index: Chunk index

        Returns:
            Unique document ID
        """
        content = f"{source}_{index}"
        return hashlib.sha256(content.encode()).hexdigest()[:16]
