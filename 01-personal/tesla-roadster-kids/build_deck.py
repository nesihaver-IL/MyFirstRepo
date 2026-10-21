"""Build the Hebrew (RTL) kids' deck about the Tesla Roadster demo (15 Oct 2026): tesla-roadster-kids.pptx.

Run from this folder:  .venv/bin/python build_deck.py
Optional: put a photo with a verified open license at images/musk.jpg (credit in images/CREDIT.txt).
Requires: python-pptx. The video is embedded from video/roadster-demo-animation.mp4.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
VIDEO = HERE / "video" / "roadster-demo-animation.mp4"
POSTER = HERE / "video" / "poster.png"
MUSK_PHOTO = HERE / "images" / "musk.jpg"
OUTPUT = HERE / "tesla-roadster-kids.pptx"

FONT = "Arial"  # Arial renders Hebrew in PowerPoint on Windows and macOS

# Tesla-style palette. Hex values are approximate (see README): verify against Tesla's press materials.
BG = RGBColor(0xF4, 0xF4, 0xF4)
RED = RGBColor(0xE8, 0x21, 0x27)
RED_DARK = RGBColor(0xA8, 0x10, 0x18)
GRAPHITE = RGBColor(0x17, 0x1A, 0x20)
GREY = RGBColor(0x5C, 0x5E, 0x62)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x17, 0x1A, 0x20)
LINK = RGBColor(0xB0, 0x1A, 0x20)

CARD_COLORS = [RED, GRAPHITE, GREY, RED_DARK]


def style_paragraph(paragraph, size: int, color: RGBColor, bold: bool = False) -> None:
    """Right-to-left paragraph with Hebrew language tagging."""
    paragraph.alignment = PP_ALIGN.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    for run in paragraph.runs:
        font = run.font
        font.name = FONT
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rPr.set("lang", "he-IL")
        rPr.set("altLang", "en-US")


def add_text(slide, left, top, width, height, lines, size=24, color=INK, bold=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        style_paragraph(p, size, color, bold)
        p.space_after = Pt(10)
    return box


def add_link(slide, left, top, width, height, label: str, url: str, size: int = 20) -> None:
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = label
    run.hyperlink.address = url
    style_paragraph(p, size, LINK, bold=True)


def add_card(slide, left, top, width, height, title, body, color, text_color=WHITE, icon=None):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.12
    card.fill.solid()
    card.fill.fore_color.rgb = color
    card.line.fill.background()
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.25)
    if icon is not None:
        tf.margin_left = Inches(1.25)
        add_icon(slide, icon, left + Inches(0.3), top + (height - Inches(0.7)) // 2, Inches(0.7), color)
    p = tf.paragraphs[0]
    p.text = title
    style_paragraph(p, 26, text_color, bold=True)
    p2 = tf.add_paragraph()
    p2.text = body
    style_paragraph(p2, 20, text_color)
    return card


def add_icon(slide, shape_type, left, top, size, card_color) -> None:
    """Native-shape icon in a white disc on the card (no logos, no third-party artwork)."""
    disc = slide.shapes.add_shape(MSO_SHAPE.OVAL, left, top, size, size)
    disc.fill.solid()
    disc.fill.fore_color.rgb = WHITE
    disc.line.fill.background()
    inner_size = int(size * 0.55)
    inner = slide.shapes.add_shape(shape_type, left + (size - inner_size) // 2,
                                   top + (size - inner_size) // 2, inner_size, inner_size)
    inner.fill.solid()
    inner.fill.fore_color.rgb = card_color
    inner.line.fill.background()


def add_background(slide, prs) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, prs, text: str) -> None:
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.12))
    band.fill.solid()
    band.fill.fore_color.rgb = RED
    band.line.fill.background()
    add_text(slide, Inches(0.6), Inches(0.35), prs.slide_width - Inches(1.2), Inches(1.0),
             [text], size=40, color=GRAPHITE, bold=True)


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    return slide


def slide_title(prs) -> None:
    slide = new_slide(prs)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.35))
    band.fill.solid()
    band.fill.fore_color.rgb = RED
    band.line.fill.background()
    add_text(slide, Inches(0.8), Inches(1.5), prs.slide_width - Inches(1.6), Inches(1.6),
             ["טסלה רודסטר: הרכב שקפץ באוויר"], size=50, color=GRAPHITE, bold=True)
    add_text(slide, Inches(0.8), Inches(3.1), prs.slide_width - Inches(1.6), Inches(0.8),
             ["ההדגמה מ-15 באוקטובר 2026"], size=30, color=GREY)
    add_text(slide, Inches(0.8), Inches(3.95), prs.slide_width - Inches(1.6), Inches(1.2),
             ["האם מכונית יכולה לטפס על רמפה, להסתובב באוויר ולנחות על הגלגלים?"],
             size=30, color=RED, bold=True)
    add_text(slide, Inches(0.8), Inches(5.6), prs.slide_width - Inches(1.6), Inches(0.6),
             ["מצגת לכיתה ו׳"], size=20, color=INK)
    add_notes(slide,
              "פתיחה: שאלו את הילדים מה לדעתם מכונית צריכה כדי לקפוץ ולהסתובב באוויר.\n\n"
              "מקורות: https://www.cnn.com/2026/10/16/world/video/tesla-unveils-what-it-says-is-a-flying-car ; "
              "https://www.statesman.com/business/article/tesla-roadster-mcgregor-flying-car-musk-spacex-22480443.php")


def slide_what(prs) -> None:
    slide = new_slide(prs)
    add_title(slide, prs, "מה קרה?")
    add_text(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.8), [
        "טסלה חשפה את הדור החדש של הרודסטר, מכונית ספורט חשמלית.",
        "הרכב נסע במעלה רמפה מעוגלת, התהפך עם הגג כלפי מטה, קפץ מקצה הרמפה והסתובב באוויר.",
        "הוא נחת על הגלגלים. לפי הדיווחים, זו הייתה קפיצה קצרה ולא טיסה ממושכת.",
    ], size=26)
    add_notes(slide,
              "מקורות: https://www.cnn.com/2026/10/16/world/video/tesla-unveils-what-it-says-is-a-flying-car (קפיצה קצרה, לא טיסה ממושכת, לפי מומחה בתעופה), "
              "https://www.statesman.com/business/article/tesla-roadster-mcgregor-flying-car-musk-spacex-22480443.php (הרמפה והתהפכות).")


def slide_when_where(prs) -> None:
    slide = new_slide(prs)
    add_title(slide, prs, "מתי ואיפה?")
    add_card(slide, Inches(0.6), Inches(1.7), Inches(5.9), Inches(2.3),
             "15 באוקטובר 2026", "ההדגמה התקיימה ביום חמישי, אחרי דחייה מ-1 באוקטובר בגלל מזג אוויר",
             RED, icon=MSO_SHAPE.FOLDED_CORNER)
    add_card(slide, Inches(6.8), Inches(1.7), Inches(5.9), Inches(2.3),
             "מקגרגור, טקסס", "במתקן הבדיקות של SpaceX, חברה שמאסק מנהל גם היא",
             GRAPHITE, icon=MSO_SHAPE.TEAR)
    add_text(slide, Inches(0.8), Inches(4.5), Inches(11.7), Inches(2.2), [
        "שימו לב: המועד והמיקום מבוססים על דיווחי חדשות. לפני ההצגה, בדקו את הפרטים בכתבה המקורית.",
    ], size=22, color=GREY)
    add_notes(slide,
              "מקורות: https://driveteslacanada.ca/news/tesla-roadster-unveiling-livestream-october-15-2026/ ; "
              "https://arynews.tv/tesla-to-unveil-long-delayed-roadster-on-october-1-musk-says (דחייה מ-1 באוקטובר בגלל מזג אוויר). "
              "הפרט 'חמישי' נלקח מאותם דיווחים: 15 באוקטובר 2026 חל ביום חמישי. מיקום: מתקן SpaceX ב-McGregor, טקסס.")


def slide_who(prs) -> None:
    slide = new_slide(prs)
    add_title(slide, prs, "מי מעורב?")
    add_card(slide, Inches(0.6), Inches(1.7), Inches(3.9), Inches(2.4),
             "טסלה", "יצרנית הרכב. הרודסטר הוא רכב ספורט חשמלי שלה", RED)
    add_card(slide, Inches(4.7), Inches(1.7), Inches(3.9), Inches(2.4),
             "SpaceX", "פיתחה יחד עם טסלה את סילוני הגז הקרים לדחף", GRAPHITE)
    add_card(slide, Inches(8.8), Inches(1.7), Inches(3.9), Inches(2.4),
             "אילון מאסק", "מוביל את טסלה ואת SpaceX לפי הדיווחים. אנשים רבים רואים בו מוביל חדשנות", GREY)
    if MUSK_PHOTO.exists():
        slide.shapes.add_picture(str(MUSK_PHOTO), Inches(0.6), Inches(4.4), height=Inches(2.6))
    else:
        frame = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(4.4), Inches(3.2), Inches(2.6))
        frame.fill.solid()
        frame.fill.fore_color.rgb = WHITE
        frame.line.color.rgb = GREY
        tf = frame.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = "כאן תוצב תמונה של אילון מאסק מרישיון פתוח, עם קרדיט"
        style_paragraph(p, 16, GREY)
    add_text(slide, Inches(4.2), Inches(4.5), Inches(8.5), Inches(2.4), [
        "מאסק הוא הדמות המרכזית בטסלה וב-SpaceX, ומכאן הקשר בין שתי החברות בהדגמה.",
    ], size=24, color=INK)
    add_notes(slide,
              "מקורות לתפקידים: https://builtin.com/articles/elon-musk (סקירה כללית, לאמת). "
              "הדיווחים על תפקידיו העדכניים לא תואמים זה את זה, ולכן כדאי לאמת מול אתר טסלה ו-SpaceX לפני ההצגה. "
              "תמונה: יש להוסיף קובץ images/musk.jpg רק אחרי שבדקתם את הרישיון (למשל Wikimedia Commons), ולרשום את הקרדיט ב-images/CREDIT.txt.")


def slide_how(prs) -> None:
    slide = new_slide(prs)
    add_title(slide, prs, "איך זה עובד?")
    steps = [
        ("מאווררים מתחת לרכב", "יניקה שמחזיקה את הרכב על הרמפה, לפי טסלה"),
        ("סילוני גז קרים", "דחיפה קטנה שעוזרת לקפוץ ולהסתובב"),
        ("עדיין לא מפורט", "מספר הסילונים מדווח שונה בין מקורות"),
    ]
    top = Inches(1.7)
    for i, (title, body) in enumerate(steps):
        add_card(slide, Inches(0.6), top + Inches(1.5) * i, Inches(5.2), Inches(1.3),
                 title, body, CARD_COLORS[i])
    video_left, video_top = Inches(6.4), Inches(1.7)
    video_width = Inches(6.3)
    video_height = Inches(6.3 * 9 / 16)
    poster = POSTER if POSTER.exists() else None
    if VIDEO.exists():
        slide.shapes.add_movie(str(VIDEO), video_left, video_top, video_width, video_height,
                               poster_frame_image=str(poster) if poster else None)
    add_text(slide, video_left, video_top + video_height + Inches(0.15), video_width, Inches(0.5),
             ["אנימציה שממחישה את ההדגמה. לא צילום מקורי"], size=18, color=GREY)
    add_notes(slide,
              "הסרטון הוא אנימציה המבוססת על תיאורי ההדגמה, לא צילום שלה.\n"
              "מאווררים וגז קר: לפי טסלה, כמתואר ב-https://www.statesman.com/business/article/tesla-roadster-mcgregor-flying-car-musk-spacex-22480443.php ; "
              "https://aviationweek.com/aerospace/emerging-technologies/tesla-demos-flying-roadster-powered-spacex-cold-gas-thrusters .\n"
              "מספר הפתחים: מוסק דיבר על 10 פתחים, דיווח אחר על 4 מאחור ועוד 6 בגרסת Rocket. לכן לא פירטנו את המספר.\n"
              "זמן 0-60 של 0.81 שנייה הוא מספר של טסלה ו-Aviation Week. 'פחות משנייה' הוא של מוסק. "
              "אלה טענות של החברה, ולכן לא הופיעו בשקף.")


def slide_why(prs) -> None:
    slide = new_slide(prs)
    add_title(slide, prs, "למה זה חשוב, ומה אנחנו חושבים?")
    add_text(slide, Inches(0.6), Inches(1.5), Inches(12.1), Inches(2.6), [
        "למה זה חשוב? ההדגמה מראה איך חברות משתפות פעולה כדי לנסות דברים שלא נעשו לפני, גם כשזה נראה מסוכן.",
        "מה אנחנו חושבים? אילון מאסק הוא מוביל חכם וחדשן, שמביא פיתוחים חדשים לעולם, כולל רכב חשמלי. "
        "אבל ההדגמה מראה שהחלום לרכב שמתעופף עדיין רחוק מהמוצר שאפשר לקנות.",
    ], size=24)
    add_card(slide, Inches(0.6), Inches(4.7), Inches(12.1), Inches(1.9),
             "שאלה לדיון", "מה דעתכם? האם חדשנות צריכה להיות מהירה, או בטוחה ומדויקת?", RED_DARK)
    add_notes(slide,
              "הדעה בשקף היא דעה קבוצתית, לא של כל ילד בנפרד. יש לנסח אותה מחדש עם הקבוצה לפני ההצגה.\n"
              "הטענה 'מוביל חכם וחדשן' היא דעה. הטענה 'מביא פיתוחים לעולם' נתמכת במידה רבה על ידי המקורות על טסלה ו-SpaceX.\n"
              "טסלה לא הייתה הרכב החשמלי הראשון, ולכן לא כתבנו שהיא הראשונה. "
              "הטענה 'עשיר מאוד' הוסרה כי לא אומתה במקור.")


def slide_links(prs) -> None:
    slide = new_slide(prs)
    add_title(slide, prs, "הסרטון והכתבה בעברית")
    add_card(slide, Inches(0.6), Inches(1.6), Inches(5.9), Inches(2.6),
             "הסרטון כפי שפורסם באתר חדשות", "CNN. לחצו על הקישור לצפייה בסרטון הרשמי של ההדגמה, לא מוטמע כאן",
             GRAPHITE)
    add_link(slide, Inches(0.9), Inches(3.6), Inches(5.4), Inches(0.6),
             "לצפייה בסרטון ב-CNN", "https://www.cnn.com/2026/10/16/world/video/tesla-unveils-what-it-says-is-a-flying-car")
    add_card(slide, Inches(6.8), Inches(1.6), Inches(5.9), Inches(2.6),
             "כתבה בעברית", "i24NEWS: הכתבה מתארת את ההדגמה ואת טענות מאסק על התמרונים",
             RED)
    add_link(slide, Inches(7.1), Inches(3.6), Inches(5.4), Inches(0.6),
             "לקריאת הכתבה ב-i24NEWS", "https://www.i24news.tv/he/news/science-and-technology/artc-28d73fdd")
    add_text(slide, Inches(0.6), Inches(4.6), Inches(12.1), Inches(2.4), [
        "מה הכתבה אומרת, במילים שלנו: טסלה חשפה את הדור החדש של הרודסטר. מאסק טען שהרכב יכול להתרומם ולבצע תמרונים. "
        "ההדגמה נערכה במתקן של SpaceX בטקסס, לא הוצגה טיסה ממושכת, ולא נמסר תאריך מסירה ללקוחות.",
    ], size=22)
    add_notes(slide,
              "קישורים: https://www.cnn.com/2026/10/16/world/video/tesla-unveils-what-it-says-is-a-flying-car (סרטון), "
              "https://www.i24news.tv/he/news/science-and-technology/artc-28d73fdd (כתבה בעברית). "
              "כתבה נוספת בעברית: https://www.ice.co.il/auto/news/article/1133858 . "
              "לפני ההצגה, פתחו את הקישורים ובדקו שהם עדיין פעילים. הסרטון לא מוטמע כי הוא חומר מוגן זכויות.")


def slide_formats(prs) -> None:
    slide = new_slide(prs)
    add_title(slide, prs, "הסיפור בכל דרך")
    add_card(slide, Inches(0.6), Inches(1.6), Inches(2.9), Inches(2.3), "סרטון", "האנימציה בשקף 5", RED,
             icon=MSO_SHAPE.ISOSCELES_TRIANGLE)
    add_card(slide, Inches(3.7), Inches(1.6), Inches(2.9), Inches(2.3), "תמונה", "תמונת מאסק בשקף 4", GRAPHITE,
             icon=MSO_SHAPE.FRAME)
    add_card(slide, Inches(6.8), Inches(1.6), Inches(2.9), Inches(2.3), "שיר", "מילים בצד ימין", GREY,
             icon=MSO_SHAPE.WAVE)
    add_card(slide, Inches(9.9), Inches(1.6), Inches(2.9), Inches(2.3), "שאלה", "שאלה לדיון למטה", RED_DARK,
             icon=MSO_SHAPE.DONUT)
    add_text(slide, Inches(9.2), Inches(4.2), Inches(3.6), Inches(3.0), [
        "רכב אדום עולה על הרמפה,",
        "מסתובב באוויר וחוזר לארץ.",
        "טסלה וחלל חולמים יחד,",
        "וכל ניסיון מלמד משהו חדש.",
    ], size=22, color=INK, bold=True)
    add_text(slide, Inches(0.6), Inches(4.3), Inches(8.2), Inches(2.6), [
        "שיר קצר, מילים מקוריות של הכיתה. שירו את הבית על המנגינה שאתם אוהבים.",
        "שאלה לדיון: מה הייתם שואלים את מאסק אם הייתם פוגשים אותו?",
    ], size=24)
    add_notes(slide,
              "שאלה 6 מבקשת להציג את האירוע בכל דרך. בשקף הזה הצגנו: סרטון, תמונה, שיר קצר ושאלה לדיון.\n"
              "השיר נכתב לשקף הזה, ואינו מבוסס על חומר מוגן זכויות. ניתן לשנות אותו לפי הרצון.")


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_title(prs)
    slide_what(prs)
    slide_when_where(prs)
    slide_who(prs)
    slide_how(prs)
    slide_why(prs)
    slide_links(prs)
    slide_formats(prs)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"wrote {build()}")
