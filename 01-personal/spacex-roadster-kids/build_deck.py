"""Build the Hebrew (RTL) kids' deck about the Falcon Heavy demo flight (Tesla Roadster): falcon-heavy-kids.pptx.

Run from this folder:  .venv/bin/python build_deck.py
Requires: python-pptx. The video is embedded from video/falcon-heavy-animation.mp4.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
VIDEO = HERE / "video" / "falcon-heavy-animation.mp4"
POSTER = HERE / "video" / "poster.png"
OUTPUT = HERE / "falcon-heavy-kids.pptx"

FONT = "Arial"  # Arial renders Hebrew in PowerPoint on Windows and macOS

BG = RGBColor(0xFF, 0xF8, 0xE7)
NAVY = RGBColor(0x1B, 0x3A, 0x6B)
ORANGE = RGBColor(0xFF, 0x8C, 0x28)
YELLOW = RGBColor(0xFF, 0xD5, 0x4F)
GREEN = RGBColor(0x3B, 0xA5, 0x5D)
PINK = RGBColor(0xE8, 0x5D, 0x9B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x22, 0x33)

CARD_COLORS = [ORANGE, GREEN, PINK, NAVY]


def style_paragraph(paragraph, size: int, color: RGBColor, bold: bool = False) -> None:
    """Right-to-left paragraph with Hebrew language tagging."""
    paragraph.alignment = PP_ALIGN.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    for run in paragraph.runs:
        font = run.font
        font.name = FONT
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rPr.set("lang", "he-IL")
        rPr.set("altLang", "en-US")


def add_text(slide, left, top, width, height, lines, size=24, color=INK, bold=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        style_paragraph(p, size, color, bold)
        p.space_after = Pt(10)
    return box


def add_card(slide, left, top, width, height, title, body, color, text_color=WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.12
    card.fill.solid()
    card.fill.fore_color.rgb = color
    card.line.fill.background()
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = title
    style_paragraph(p, 26, text_color, bold=True)
    p2 = tf.add_paragraph()
    p2.text = body
    style_paragraph(p2, 20, text_color)
    return card


def add_background(slide, prs) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, prs, text: str) -> None:
    add_text(slide, Inches(0.6), Inches(0.35), prs.slide_width - Inches(1.2), Inches(1.0),
             [text], size=40, color=NAVY, bold=True)


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def slide_title(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.35))
    band.fill.solid()
    band.fill.fore_color.rgb = ORANGE
    band.line.fill.background()
    add_text(slide, Inches(0.8), Inches(1.6), prs.slide_width - Inches(1.6), Inches(1.6),
             ["מכונית בחלל!"], size=60, color=NAVY, bold=True)
    add_text(slide, Inches(0.8), Inches(3.0), prs.slide_width - Inches(1.6), Inches(0.8),
             ["הטיסה הראשונה של Falcon Heavy, 2018"], size=30, color=NAVY)
    add_text(slide, Inches(0.8), Inches(3.9), prs.slide_width - Inches(1.6), Inches(1.2),
             ["למה שמישהו ישלח מכונית לחלל?"], size=30, color=ORANGE, bold=True)
    add_text(slide, Inches(0.8), Inches(5.6), prs.slide_width - Inches(1.6), Inches(0.6),
             ["מצגת לכיתה ו׳  |  SpaceX"], size=20, color=INK)
    add_notes(slide,
              "פתיחה: שאלו את הילדים למה לדעתם שולחים מכונית לחלל, לפני שמספרים להם את התשובה.\n\n"
              "מקורות לבדיקה: https://www.pressreader.com/usa/albuquerque-journal/20180207/281595240984804 ; "
              "https://planeandpilotmag.com/spacex-launches-falcon-heavy-rocket/")


def slide_what(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "מה שוגר לחלל?")
    add_text(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.5), [
        "הרקטה: Falcon Heavy, שלושה בוסטרים שעובדים ביחד, עם כ-27 מנועים.",
        "באותו זמן זו הייתה הרקטה החזקה ביותר שפועלת בעולם.",
        "המטען: מכונית Tesla Roadster אדומה, עם בובה בשם Starman בחליפת חלל.",
        "הטיסה יצאה ב-6 בפברואר 2018 מאתר השיגור 39A בפלורידה.",
    ], size=26)
    add_notes(slide,
              "מספר המנועים (כ-27) והכינוי 'החזקה ביותר שפועלת' מבוססים על מקור משני בלבד: "
              "https://weatherboy.com/?p=4675 . יש לאמת מול אתר SpaceX לפני ההצגה.\n"
              "תאריך השיגור ואתר השיגור: https://planeandpilotmag.com/spacex-launches-falcon-heavy-rocket/")


def slide_process(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "איך הכינו את הטיסה?")
    steps = [
        ("משקל במקום מטען", "במקום מטען רגיל בחרו מכונית"),
        ("Starman בחליפה", "בובה בחליפת חלל שתוכננה ב-SpaceX"),
        ("שידור חי לעולם", "מיליוני אנשים צפו בטיסה בשידור חי"),
    ]
    top = Inches(1.7)
    for i, (title, body) in enumerate(steps):
        add_card(slide, Inches(0.6), top + Inches(1.5) * i, Inches(5.2), Inches(1.3),
                 title, body, CARD_COLORS[i])
    video_left, video_top = Inches(6.4), Inches(1.7)
    video_width = Inches(6.3)
    video_height = Inches(6.3 * 9 / 16)
    poster = POSTER if POSTER.exists() else None
    if VIDEO.exists():
        slide.shapes.add_movie(str(VIDEO), video_left, video_top, video_width, video_height,
                               poster_frame_image=str(poster) if poster else None)
    add_text(slide, video_left, video_top + video_height + Inches(0.15), video_width, Inches(0.5),
             ["סרטון קצר: שלוש רקטות, שני בוסטרים שחוזרים, והמכונית בחלל"], size=18, color=INK)
    add_notes(slide,
              "הציגו את הסרטון (כ-36 שניות). הסרטון הוא אנימציה המבוססת על מה שהמקורות מתארים, לא צילום אמיתי.\n"
              "מקורות: https://www.nbcchicago.com/news/national-international/spacex-falcon-heavy-first-launch/2041756/ "
              "(המשקל והמכונית, וציטוט של מאסק שמשקל פלדה או בטון 'נשמע משעמם').\n"
              "שידור חי: אותו מקור מתאר מיליוני צופים ברשת.")


def slide_result(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "מה קרה בטיסה?")
    cards = [
        ("שני בוסטרים נחתו", "הבוסטרים הצדדיים חזרו ונחתו על הקרקע, ליד השיגור"),
        ("המרכזי החטיא", "הבוסטר המרכזי פספס את הספינה הימית ונפל לים"),
        ("שימוש חוזר", "בוסטרים שחוזרים אפשר להשתמש שוב, וזה חוסך כסף"),
        ("המכונית בחלל", "היא הגיעה למסלול מעבר למאדים, ולא בדיוק למה שתוכנן"),
    ]
    col_w, row_h = Inches(5.9), Inches(2.2)
    for i, (title, body) in enumerate(cards):
        col, row = i % 2, i // 2
        left = Inches(0.6) + col * (col_w + Inches(0.4))
        top = Inches(1.7) + row * (row_h + Inches(0.3))
        add_card(slide, left, top, col_w, row_h, title, body, CARD_COLORS[i])
    add_notes(slide,
              "שני הבוסטרים הצדדיים נחתו במתחמים Landing Zone 1 ו-2 בקייפ קנדי: "
              "https://thedailyworld.com/?p=16224 , https://lunarcabin.com/Rocket_Launches/falconheavydemo_1.html .\n"
              "הבוסטר המרכזי התכוון לספינה הימית 'Of Course I Still Love You' ופספס בכ-100 מטר: "
              "https://thedailyworld.com/?p=16224 .\n"
              "המסלול של המכונית: https://orbitcodex.com/operations/demo-flight-1-maiden-flight-falcon-heavy .\n"
              "לא אומת ולכן לא הופיע בשקף: מוזיקה שהתנגנה במכונית, ולוח מעגל עם טקסט על המכונית.")


def slide_empowerment(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "גם מהחסרונות לומדים!")
    messages = [
        ("שאלו שאלות", "כל שאלה טובה, גם אם נראית קטנה"),
        ("כשלון הוא שיעור", "גם הבוסטר שפספס למד את המהנדסים הרבה"),
        ("חלמו בגדול", "חלום גדול מתחיל מצעד קטן אחד"),
    ]
    for i, (title, body) in enumerate(messages):
        add_card(slide, Inches(0.6) + i * Inches(4.1), Inches(1.7), Inches(3.9), Inches(2.4),
                 title, body, CARD_COLORS[i])
    add_card(slide, Inches(0.6), Inches(4.6), prs.slide_width - Inches(1.2), Inches(1.9),
             "שאלה לדיון", "אם הייתם שולחים משהו לחלל, מה זה היה ולמה?", YELLOW, text_color=INK)
    add_notes(slide,
              "סיום: תנו לכל ילד לענות בקצרה. אפשר לחלק כרטיסי כתיבה ולאסוף את התשובות לשיעור הבא.")


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_title(prs)
    slide_what(prs)
    slide_process(prs)
    slide_result(prs)
    slide_empowerment(prs)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"wrote {build()}")
