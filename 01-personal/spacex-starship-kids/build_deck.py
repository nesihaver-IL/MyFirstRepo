"""Build the Hebrew (RTL) kids' deck about Starship: starship-kids.pptx.

Run from this folder:  .venv/bin/python build_deck.py
Requires: python-pptx. The video is embedded from video/starship-animation.mp4.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).parent
VIDEO = HERE / "video" / "starship-animation.mp4"
POSTER = HERE / "video" / "poster.png"
OUTPUT = HERE / "starship-kids.pptx"

FONT = "Arial"  # Arial renders Hebrew in PowerPoint on Windows and macOS

BG = RGBColor(0xFF, 0xF8, 0xE7)
NAVY = RGBColor(0x1B, 0x3A, 0x6B)
ORANGE = RGBColor(0xFF, 0x8C, 0x28)
YELLOW = RGBColor(0xFF, 0xD5, 0x4F)
GREEN = RGBColor(0x3B, 0xA5, 0x5D)
PINK = RGBColor(0xE8, 0x5D, 0x9B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
INK = RGBColor(0x22, 0x22, 0x33)

CARD_COLORS = [ORANGE, GREEN, PINK, NAVY]


def style_paragraph(paragraph, size: int, color: RGBColor, bold: bool = False) -> None:
    """Right-to-left paragraph with Hebrew language tagging."""
    paragraph.alignment = PP_ALIGN.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    pPr.set("rtl", "1")
    for run in paragraph.runs:
        font = run.font
        font.name = FONT
        font.size = Pt(size)
        font.bold = bold
        font.color.rgb = color
        rPr = run._r.get_or_add_rPr()
        rPr.set("lang", "he-IL")
        rPr.set("altLang", "en-US")


def add_text(slide, left, top, width, height, lines, size=24, color=INK, bold=False, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        style_paragraph(p, size, color, bold)
        p.space_after = Pt(10)
    return box


def add_card(slide, left, top, width, height, title, body, color, text_color=WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.adjustments[0] = 0.12
    card.fill.solid()
    card.fill.fore_color.rgb = color
    card.line.fill.background()
    tf = card.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.25)
    p = tf.paragraphs[0]
    p.text = title
    style_paragraph(p, 26, text_color, bold=True)
    p2 = tf.add_paragraph()
    p2.text = body
    style_paragraph(p2, 20, text_color)
    return card


def add_background(slide, prs) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG


def add_title(slide, prs, text: str) -> None:
    add_text(slide, Inches(0.6), Inches(0.35), prs.slide_width - Inches(1.2), Inches(1.0),
             [text], size=40, color=NAVY, bold=True)


def add_notes(slide, text: str) -> None:
    slide.notes_slide.notes_text_frame.text = text


def slide_title(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.35))
    band.fill.solid()
    band.fill.fore_color.rgb = ORANGE
    band.line.fill.background()
    add_text(slide, Inches(0.8), Inches(1.6), prs.slide_width - Inches(1.6), Inches(1.6),
             ["Starship: הרקטה שחוזרת הביתה"], size=54, color=NAVY, bold=True)
    add_text(slide, Inches(0.8), Inches(3.3), prs.slide_width - Inches(1.6), Inches(1.2),
             ["איך נראית רקטה שאפשר להשתמש בה שוב ושוב?"], size=30, color=ORANGE, bold=True)
    add_text(slide, Inches(0.8), Inches(5.6), prs.slide_width - Inches(1.6), Inches(0.6),
             ["מצגת לכיתה ו׳  |  על SpaceX ועל Starship"], size=20, color=INK)
    add_notes(slide, "פתיחה: שאלו את הילדים מה לדעתם יכולה לעשות רקטה שחוזרת לכדור הארץ.")


def slide_product(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "מה זה Starship?")
    add_text(slide, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.5), [
        "Starship היא רקטה ענקית שפותחה על ידי SpaceX.",
        "היא בנויה משני חלקים: הבוסטר (המאיץ) שדוחף אותה בהתחלה, והחללית שממשיכה לחלל.",
        "המטרה: לקחת אנשים ומטענים לחלל, ולחזור ולטוס שוב ושוב.",
        "היא אחת הרקטות הגדולות והחזקות שנבנו אי פעם.",
    ], size=26)
    add_notes(slide,
              "הערה למציג: הגובה והעוצמה של הרקטה משתנים בין גרסאות וטיסות. "
              "לפני ההצגה, בדקו את המספרים העדכניים באתר SpaceX (spacex.com) ובדפי הטיסות של Starship. "
              "אין לי אישור עדכני למספרים מדויקים.")


def slide_process(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "איך בונים רקטה כזאת?")
    steps = [
        ("1. בונים אב-טיפוס", "מתחילים בגרסה קטנה וניסיונית"),
        ("2. טסים ולומדים", "כל טיסת ניסוי נותנת מידע חדש"),
        ("3. נכשלים? זה חלק מהדרך", "כשלון הוא שיעור, לא סוף"),
    ]
    top = Inches(1.7)
    for i, (title, body) in enumerate(steps):
        add_card(slide, Inches(0.6), top + Inches(1.5) * i, Inches(5.2), Inches(1.3),
                 title, body, CARD_COLORS[i])
    video_left, video_top = Inches(6.4), Inches(1.7)
    video_width = Inches(6.3)
    video_height = Inches(6.3 * 9 / 16)
    poster = POSTER if POSTER.exists() else None
    if VIDEO.exists():
        slide.shapes.add_movie(str(VIDEO), video_left, video_top, video_width, video_height,
                               poster_frame_image=str(poster) if poster else None)
    add_text(slide, video_left, video_top + video_height + Inches(0.15), video_width, Inches(0.5),
             ["סרטון קצר: הרקטה עולה לחלל"], size=18, color=INK)
    add_notes(slide,
              "הציגו את הסרטון (כ-36 שניות). אחרי הסרטון: מה הילדים ראו? "
              "שאלו מה היה קורה אם הטיסה הראשונה הייתה מצליחה מיד.")


def slide_technology(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "איך זה עובד?")
    cards = [
        ("מנועי Raptor", "שורפים מתאן וחמצן נוזלי כדי לייצר דחף"),
        ("גוף מפלדה", "פלדה עמידה לחום, וחזקה יחסית למשקל שלה"),
        ("שימוש חוזר", "הבוסטר חוזר ונתפס באוויר, במקום להישרף בים"),
        ("תפיסה בזרועות", "מגדל השיגור תופס את הבוסטר עם זרועות ענק"),
    ]
    col_w, row_h = Inches(5.9), Inches(2.2)
    for i, (title, body) in enumerate(cards):
        col, row = i % 2, i // 2
        left = Inches(0.6) + col * (col_w + Inches(0.4))
        top = Inches(1.7) + row * (row_h + Inches(0.3))
        add_card(slide, left, top, col_w, row_h, title, body, CARD_COLORS[i])
    add_notes(slide,
              "הסבירו בשפה פשוטה: מתאן הוא גז שמשמש גם כדלק, וחמצן נוזלי הוא החמצן שנותן לדלק לבעור. "
              "הזרועות תופסות את הבוסטר בדומה לאיך שמיישרים כביסה על חבל. "
              "הפרטים הטכניים משתנים בין גרסאות - בדקו לפני ההצגה.")


def slide_empowerment(prs) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_background(slide, prs)
    add_title(slide, prs, "גם אתם יכולים להמריא!")
    messages = [
        ("שאלו שאלות", "כל שאלה טובה, גם אם נראית קטנה"),
        ("אל תפחדו לטעות", "גם מהנדסי SpaceX למדו מניסיונות שנכשלו"),
        ("חלמו בגדול, תכננו צעד צעד", "חלום גדול מתחיל מצעד קטן אחד"),
    ]
    for i, (title, body) in enumerate(messages):
        add_card(slide, Inches(0.6) + i * Inches(4.1), Inches(1.7), Inches(3.9), Inches(2.4),
                 title, body, CARD_COLORS[i])
    add_card(slide, Inches(0.6), Inches(4.6), prs.slide_width - Inches(1.2), Inches(1.9),
             "שאלה לדיון", "אם הייתם בונים רקטה, מה הייתם רוצים שתעשה?", YELLOW, text_color=INK)
    add_notes(slide,
              "סיום: תנו לכל ילד לענות בקצרה. אפשר לחלק כרטיסי כתיבה ולאסוף את התשובות לשיעור הבא.")


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_title(prs)
    slide_product(prs)
    slide_process(prs)
    slide_technology(prs)
    slide_empowerment(prs)
    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(f"wrote {build()}")
